"""Fetch contract sources C1 to C12 and extract plain text.

Raw downloads and extracted text stay OUTSIDE the repo (default /tmp/raw/contracts),
because filed sources contain real people's data. Only the manifest is written
into the repo: data/sources/contracts/manifest.json.

Usage: python scripts/sources/contracts/fetch.py [--raw-dir /tmp/raw/contracts] [--only C8,C9]

Each source is fetched from its `fetch` list when given, otherwise from `url`.
Publication pages (gov.uk, gov.wales, National Archives) link to the document as
an attachment: the best attachment (DOCX, then ODT, then PDF) is downloaded and
extracted; a page without attachments is extracted as HTML.
"""

import argparse
import datetime as dt
import hashlib
import io
import json
import re
import sys
import zipfile
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests

sys.path.insert(0, str(Path(__file__).resolve().parent))
from sources import SOURCES  # noqa: E402

REPO = Path(__file__).resolve().parents[3]
MANIFEST = REPO / "data" / "sources" / "contracts" / "manifest.json"
UA = "night-shift-hackathon dataset build (github.com/BabyBoss45/night-shift)"
DOC_EXT = (".docx", ".odt", ".pdf")


def get(url):
    r = requests.get(url, headers={"User-Agent": UA}, timeout=60)
    r.raise_for_status()
    return r


def tidy(text):
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"


def md_to_text(md):
    out = []
    for ln in md.split("\n"):
        ln = re.sub(r"</?(label|span)[^>]*>", "", ln)
        ln = ln.replace("**", "")
        ln = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", ln)
        ln = re.sub(r"^#+\s*", "", ln)
        ln = re.sub(r"^- (\[[ x]\])\s+", r"\1 ", ln)
        if ln.startswith("|"):
            if re.fullmatch(r"[|:\- ]+", ln):
                continue
            cells = [c.strip() for c in ln.strip().strip("|").split("|")]
            ln = " | ".join(cells)
        out.append(ln)
    return tidy("\n".join(out))


def html_to_text(raw):
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(raw, "lxml")
    for tag in soup(["script", "style", "noscript", "nav", "header", "footer", "form"]):
        tag.decompose()
    root = soup.find("main") or soup.find("article") or soup.body or soup
    for br in root.find_all("br"):
        br.replace_with("\n")
    blocks = []
    for el in root.find_all(["p", "h1", "h2", "h3", "h4", "h5", "li", "tr", "pre", "div"]):
        if el.name == "div" and el.find(["p", "div", "table", "li", "h1", "h2", "h3"]):
            continue
        if el.name in ("p", "li") and el.find_parent(["li", "tr"]) is not None:
            continue
        if el.name == "tr":
            txt = " | ".join(c.get_text(" ", strip=True) for c in el.find_all(["td", "th"]))
        else:
            txt = el.get_text(" ", strip=el.name != "pre")
        if txt.strip():
            blocks.append(txt)
    return tidy("\n\n".join(blocks) if blocks else root.get_text("\n"))


def pdf_to_text(data):
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return tidy("\n\n".join(p.extract_text() or "" for p in pdf.pages))


def docx_to_text(data):
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return tidy("\n\n".join(parts))


def odt_to_text(data):
    from lxml import etree

    xml = zipfile.ZipFile(io.BytesIO(data)).read("content.xml")
    ns = {"text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0"}
    root = etree.fromstring(xml)
    paras = ["".join(el.itertext()) for el in root.iterfind(".//text:p", ns)]
    return tidy("\n\n".join(paras))


def sec_txt_to_text(raw):
    # EDGAR .txt exhibits may wrap the text in SGML tags such as <DOCUMENT>, <TEXT>.
    return tidy(re.sub(r"^<[^>\n]+>$", "", raw, flags=re.M))


def extract(url, resp):
    path = urlparse(url).path.lower()
    ctype = resp.headers.get("content-type", "")
    if path.endswith(".md"):
        return md_to_text(resp.text)
    if path.endswith(".pdf") or "pdf" in ctype:
        return pdf_to_text(resp.content)
    if path.endswith(".docx") or "wordprocessingml" in ctype:
        return docx_to_text(resp.content)
    if path.endswith(".odt") or "opendocument.text" in ctype:
        return odt_to_text(resp.content)
    if path.endswith(".txt"):
        return sec_txt_to_text(resp.text)
    return html_to_text(resp.text)


def pick_attachment(page_url, html, hint):
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "lxml")
    found = []
    for a in soup.find_all("a", href=True):
        href = urljoin(page_url, a["href"])
        p = urlparse(href).path.lower()
        if p.endswith(DOC_EXT):
            label = (a.get_text(" ", strip=True) + " " + p).lower()
            if hint and hint not in label:
                continue
            found.append((DOC_EXT.index(Path(p).suffix), href))
    return sorted(found)[0][1] if found else None


def fetch_source(src, raw_dir):
    urls = src.get("fetch")
    notes = []
    if not urls:
        page = get(src["url"])
        att = None
        if "html" in page.headers.get("content-type", ""):
            att = pick_attachment(src["url"], page.text, src.get("attachment_hint"))
        urls = [att] if att else [src["url"]]
        if att:
            notes.append(f"attachment: {att}")
    texts = []
    for i, u in enumerate(urls):
        r = get(u)
        name = f"{src['source_id']}-{i}{Path(urlparse(u).path).suffix or '.html'}"
        (raw_dir / name).write_bytes(r.content)
        texts.append(extract(u, r))
    return "\n".join(texts), urls, notes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw-dir", default="/tmp/raw/contracts")
    ap.add_argument("--only", default="")
    args = ap.parse_args()
    raw_dir = Path(args.raw_dir)
    (raw_dir / "text").mkdir(parents=True, exist_ok=True)
    only = set(filter(None, args.only.split(",")))

    old = {}
    if MANIFEST.exists():
        old = {e["source_id"]: e for e in json.loads(MANIFEST.read_text())["sources"]}
    entries = []
    for src in SOURCES:
        sid = src["source_id"]
        if only and sid not in only:
            entries.append(old.get(sid, {"source_id": sid}))
            continue
        entry = {
            "source_id": sid,
            "title": src["title"],
            "url": src["url"],
            "licence": src["licence"],
            "kind": src["kind"],
            "subtype": src["subtype"],
        }
        try:
            text, used, notes = fetch_source(src, raw_dir)
        except Exception as e:  # record and continue with the other sources
            prev = old.get(sid, {})
            host = urlparse((src.get("fetch") or [src["url"]])[0]).hostname
            if isinstance(e, requests.exceptions.ProxyError):
                why = f"blocked: the build sandbox's egress proxy refused CONNECT to {host} (403)"
            else:
                why = f"fetch failed for {host}: {type(e).__name__}: {str(e)[:160]}"
            entry.update(
                fetched=prev.get("fetched", False),
                fetched_at=prev.get("fetched_at"),
                fetched_from=prev.get("fetched_from"),
                sha256=prev.get("sha256"),
                notes=why,
            )
            print(f"{sid}: FAILED {entry['notes']}", file=sys.stderr)
            entries.append(entry)
            continue
        (raw_dir / "text" / f"{sid}.txt").write_text(text, encoding="utf-8")
        entry.update(
            fetched=True,
            fetched_at=dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat(),
            fetched_from=used,
            sha256=hashlib.sha256(text.encode("utf-8")).hexdigest(),
            notes="; ".join(filter(None, [src.get("note", "")] + notes)),
        )
        print(f"{sid}: ok, {len(text)} chars -> {raw_dir / 'text' / (sid + '.txt')}")
        entries.append(entry)
    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(json.dumps({"sources": entries}, indent=2, ensure_ascii=False) + "\n")


if __name__ == "__main__":
    main()
